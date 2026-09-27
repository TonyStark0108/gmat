"""CR bank: links only (GMAT Club threads), grouped into sets.

Rows 1-488: GMAT Club sets, label in column C at each set's first row (the reserve).
Rows 489+: LSAT sections, label in columns D/E (+ PrepTest name in G/H) (the default).
No column holds official answers, so CR sets use the review's fallback path.
"""
import re
from collections import Counter

import openpyxl

import config

LSAT_FIRST_ROW = 489


def looks_like_answers(values):
    vals = [str(v).strip().upper() for v in values if v not in (None, "")]
    return len(vals) > 50 and sum(1 for v in vals if re.fullmatch(r"[A-E]", v)) / len(vals) > 0.9


def run():
    ws = openpyxl.load_workbook(config.CR_XLSX).active
    issues = []
    # Is there an answer column anywhere?
    answer_cols = [c for c in range(1, ws.max_column + 1)
                   if looks_like_answers(ws.cell(r, c).value for r in range(1, ws.max_row + 1))]

    reserve, lsat, cur = [], [], None
    dropped = []
    for r in range(1, ws.max_row + 1):
        num, url = ws.cell(r, 1).value, ws.cell(r, 2).value
        link = ws.cell(r, 2).hyperlink.target if ws.cell(r, 2).hyperlink else None
        url = link or url
        if not (isinstance(url, str) and url.startswith("http")):
            issues.append(f"row {r}: no link (column B = {url!r}); skipped")
            continue
        c_label = ws.cell(r, 3).value
        if r < LSAT_FIRST_ROW:
            if c_label == "Poor quality":
                dropped.append({"row": r, "url": url, "why": "marked 'Poor quality' in column C"})
                continue
            if c_label:
                cur = {"id": f"gc-{c_label.replace('TEST', '').strip().lower().replace(' ', '')}",
                       "label": f"GMAT Club CR {c_label.title()}", "questions": []}
                reserve.append(cur)
        else:
            d_label = ws.cell(r, 4).value
            if d_label:
                sec = ws.cell(r, 5).value
                pt = ws.cell(r, 7).value
                cur = {"id": f"lsat-cr-{len(lsat) + 1:02d}",
                       "label": f"LSAT {d_label.replace('TEST', 'Test').strip()} · {sec.replace('SECTION', 'Section') if sec else ''}".strip(" ·"),
                       "preptest": pt, "questions": []}
                lsat.append(cur)
        cur["questions"].append({"n": len(cur["questions"]) + 1, "row": r, "url": url})

    for s in reserve + lsat:
        s["size"] = len(s["questions"])
        # Two questions on one stimulus share a thread; both stay, the set just opens it twice.
        dup = [u for u, k in Counter(q["url"] for q in s["questions"]).items() if k > 1]
        s["shared_threads"] = len(dup)

    stats = {
        "reserve_sets": len(reserve), "reserve_questions": sum(s["size"] for s in reserve),
        "lsat_sets": len(lsat), "lsat_questions": sum(s["size"] for s in lsat),
        "lsat_set_sizes": dict(sorted(Counter(s["size"] for s in lsat).items())),
        "reserve_set_sizes": dict(sorted(Counter(s["size"] for s in reserve).items())),
        "answer_column": answer_cols or None, "dropped": dropped,
        "unique_links": len({q["url"] for s in reserve + lsat for q in s["questions"]}),
    }
    return {"default_sets": lsat, "reserve_sets": reserve, "has_official_answers": bool(answer_cols)}, stats, issues
